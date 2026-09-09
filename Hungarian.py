# -*- coding: utf-8 -*-
"""
Created on Wed May 31 19:22:30 2023

@author: avtei
"""
import re
import os
import numpy as np
import pandas, copy
import openpyxl
from scipy.optimize import linear_sum_assignment

#https://docs.scipy.org/doc/scipy-0.18.1/reference/generated/scipy.optimize.linear_sum_assignment.html
#The student that didnt make 10 choices gets his first choice as 2 instead of 1 for doing it wrong

pandas.set_option('display.max_columns', 10)

# =============================================================================
# ADDITIONAL REQUIREMENTS
# •	Aim to give at least one student to each supervisor
# •	Make sure no staff capacities are breached
# =============================================================================

# Synonym lists used to locate columns/tables regardless of exact header
# wording, so both the flat sample CSVs and a messier multi-table xlsx export
# can be read the same way.
PROJECT_ID_SYNONYMS = ("proj num", "project num", "project no", "proj no", "project number")
CAPACITY_SYNONYMS = ("capacity", "cap")
SECTION_SYNONYMS = ("supervisor", "section", "theme", "staff")
STUDENT_ID_SYNONYMS = ("student number", "student id", "student no", "matriculation")
RANK_PATTERNS = (
    re.compile(r"^(\d+)\s*(st|nd|rd|th)?\b", re.IGNORECASE),
    re.compile(r"choice\s*(\d+)", re.IGNORECASE),
)


def _clean_header(text):
    if text is None:
        return ""
    text = re.sub(r"[^a-z0-9]+", " ", str(text).lower())
    return re.sub(r"\s+", " ", text).strip()


def _norm_id(value):
    """Canonical string form of any project/student ID, so e.g. Excel's
    58.0 and a CSV's "58" compare equal; real blanks stay NaN (never "nan")."""
    if pandas.isna(value):
        return np.nan
    if isinstance(value, str):
        value = value.strip()
        return value if value else np.nan
    if isinstance(value, (int, np.integer)):
        return str(value)
    if isinstance(value, (float, np.floating)):
        return str(int(value)) if float(value).is_integer() else str(value)
    return str(value).strip()


def _to_capacity(value):
    if value is None:
        return 1
    if isinstance(value, float) and np.isnan(value):
        return 1
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return 1
        try:
            return int(float(value))
        except ValueError:
            return 1
    return int(value)


def _resolve_column(columns, synonyms, field_name):
    for col in columns:
        if any(syn in _clean_header(col) for syn in synonyms):
            return col
    raise ValueError(f"Couldn't find a '{field_name}' column among: {list(columns)}")


def _match_columns(header_texts, synonyms):
    return [idx for idx, text in enumerate(header_texts)
            if _clean_header(text) and any(syn in _clean_header(text) for syn in synonyms)]


def _locate_table(header_texts, anchor_synonyms, field_synonyms):
    """Find an anchor column (e.g. "Student Number") in this header row, then
    the other required fields in the columns nearest to it. A sheet can have
    several tables sharing a header row with overlapping column names (e.g.
    a per-supervisor summary AND a per-project table both have "Supervisor"
    and "Capacity" columns) - picking the closest match to a distinctive
    anchor keeps the right table's columns together."""
    anchor_matches = _match_columns(header_texts, anchor_synonyms)
    if not anchor_matches:
        return None
    anchor_col = anchor_matches[0]

    field_cols = {}
    for field, synonyms in field_synonyms.items():
        matches = _match_columns(header_texts, synonyms)
        if not matches:
            return None
        field_cols[field] = min(matches, key=lambda c: abs(c - anchor_col))
    return anchor_col, field_cols


def _find_table_in_workbook(wb, anchor_synonyms, field_synonyms, max_header_rows=40):
    for ws in wb.worksheets:
        max_row = min(ws.max_row, max_header_rows)
        for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_row, values_only=True)):
            located = _locate_table(list(row), anchor_synonyms, field_synonyms)
            if located is not None:
                anchor_col, field_cols = located
                return ws, row_idx + 1, anchor_col, field_cols
    return None


def _rank_columns(ws, header_row_number, anchor_col):
    header_cells = list(next(ws.iter_rows(min_row=header_row_number, max_row=header_row_number, values_only=True)))
    ranked = []
    for idx in range(anchor_col + 1, len(header_cells)):
        cleaned = _clean_header(header_cells[idx])
        if not cleaned:
            continue
        for pattern in RANK_PATTERNS:
            match = pattern.search(cleaned)
            if match:
                ranked.append((int(match.group(1)), idx))
                break
    ranked.sort(key=lambda t: t[0])
    return ranked


def _extract_rows(ws, header_row_number, col_indexes):
    rows = []
    for row in ws.iter_rows(min_row=header_row_number + 1, max_row=ws.max_row, values_only=True):
        rows.append({name: (row[idx] if idx < len(row) else None) for name, idx in col_indexes.items()})
    return rows


class StudentProjectAssignment:
    def _load_projects_csv(self, path):
        df = pandas.read_csv(path, index_col=0, encoding="utf-8-sig")
        section_col = _resolve_column(df.columns, SECTION_SYNONYMS, "Section")
        capacity_col = _resolve_column(df.columns, CAPACITY_SYNONYMS, "Capacity")
        df = df.rename(columns={section_col: "Section", capacity_col: "Capacity"})
        return df[["Section", "Capacity"]]

    def _load_projects_xlsx(self, path):
        wb = openpyxl.load_workbook(path, data_only=True)
        found = _find_table_in_workbook(
            wb, PROJECT_ID_SYNONYMS, {"Section": SECTION_SYNONYMS, "Capacity": CAPACITY_SYNONYMS}
        )
        if found is None:
            raise ValueError(
                "Couldn't find a project table (needs a project-number column plus "
                f"Supervisor/Section and Capacity columns) in any sheet of {path!r}"
            )
        ws, header_row_number, anchor_col, field_cols = found
        col_indexes = {"ProjectID": anchor_col, **field_cols}
        records = []
        for row in _extract_rows(ws, header_row_number, col_indexes):
            pid = _norm_id(row["ProjectID"])
            if pandas.isna(pid):
                continue
            section = row["Section"]
            records.append({
                "ProjectID": pid,
                "Section": "" if section is None else str(section).strip(),
                "Capacity": _to_capacity(row["Capacity"]),
            })
        if not records:
            raise ValueError(f"Found a project table header in {ws.title!r} but no data rows under it")
        return pandas.DataFrame.from_records(records).set_index("ProjectID")

    def load_projects(self, projects):
        if not os.path.exists(projects):
            raise FileNotFoundError(f"Projects file not found: {projects}")
        ext = os.path.splitext(str(projects))[1].lower()
        if ext == ".csv":
            raw = self._load_projects_csv(projects)
        elif ext in (".xlsx", ".xls", ".xlsm"):
            raw = self._load_projects_xlsx(projects)
        else:
            raise ValueError(f"Unsupported projects file type {ext!r} (expected .csv or .xlsx)")

        raw.index = raw.index.map(_norm_id)
        if raw.index.isna().any():
            raise ValueError("Found a project row with a blank/missing project ID")
        if raw.index.duplicated().any():
            dupes = raw.index[raw.index.duplicated()].unique().tolist()
            raise ValueError(f"Duplicate project ID(s) found: {dupes}")
        raw["Capacity"] = raw["Capacity"].map(_to_capacity)
        assert (raw["Capacity"] > 0).all(), "Every project needs a Capacity > 0"

        self.projects = raw

        # Split any project with Capacity > 1 into capacity-1 sibling rows,
        # e.g. "58" (Capacity 3) -> "58", "58_1", "58_2", each Capacity 1.
        new_rows = []
        for i in self.projects.index:
            capacity = int(self.projects.at[i, "Capacity"])
            if capacity > 1:
                for addition in range(1, capacity):
                    sibling = self.projects.loc[i].copy()
                    sibling["Capacity"] = 1
                    sibling.name = f"{i}_{addition}"
                    new_rows.append(sibling)
                self.projects.at[i, "Capacity"] = 1
        if new_rows:
            self.projects = pandas.concat([self.projects, pandas.DataFrame(new_rows)])

        self.projects.index = self.projects.index.astype(str)
        self.projects = self.projects.sort_index()
        print(f"Loaded {len(self.projects)} project slots.")

    def _load_choices_csv(self, path):
        df = pandas.read_csv(path, index_col=0, encoding="utf-8-sig")
        df.columns = range(1, len(df.columns) + 1)
        return df

    def _load_choices_xlsx(self, path):
        wb = openpyxl.load_workbook(path, data_only=True)
        found = _find_table_in_workbook(wb, STUDENT_ID_SYNONYMS, {})
        if found is None:
            raise ValueError(f"Couldn't find a student choices table (needs a Student Number column) in any sheet of {path!r}")
        ws, header_row_number, anchor_col, _ = found
        ranked = _rank_columns(ws, header_row_number, anchor_col)
        if not ranked:
            raise ValueError(
                f"Found a Student Number column in {ws.title!r} but no ranked "
                "choice columns (e.g. '1st', '2nd', ...) to its right"
            )
        col_indexes = {"StudentID": anchor_col}
        col_indexes.update({str(rank_pos + 1): col for rank_pos, (_, col) in enumerate(ranked)})

        records = []
        for row in _extract_rows(ws, header_row_number, col_indexes):
            sid = _norm_id(row["StudentID"])
            if pandas.isna(sid):
                continue
            record = {"StudentID": sid}
            for rank_pos in range(1, len(ranked) + 1):
                record[str(rank_pos)] = row[str(rank_pos)]
            records.append(record)
        if not records:
            raise ValueError(f"Found a student choices header in {ws.title!r} but no data rows under it")
        return pandas.DataFrame.from_records(records).set_index("StudentID")

    def load_choices(self, choices):
        if not os.path.exists(choices):
            raise FileNotFoundError(f"Choices file not found: {choices}")
        ext = os.path.splitext(str(choices))[1].lower()
        if ext == ".csv":
            raw = self._load_choices_csv(choices)
        elif ext in (".xlsx", ".xls", ".xlsm"):
            raw = self._load_choices_xlsx(choices)
        else:
            raise ValueError(f"Unsupported choices file type {ext!r} (expected .csv or .xlsx)")

        raw.index = raw.index.map(_norm_id)
        raw = raw[~pandas.isna(raw.index)]
        if raw.index.duplicated().any():
            dupes = raw.index[raw.index.duplicated()].unique().tolist()
            raise ValueError(f"Duplicate student ID(s) found: {dupes}")

        self.choices = raw.map(_norm_id)
        self.choices.columns = [str(c) for c in self.choices.columns]
        print(f"Loaded choices for {len(self.choices)} students.")

        # Insert each capacity-sibling's ID right after its base choice, when
        # that sibling exists, so students effectively get an extra shot at
        # the same supervisor's other slot(s).
        n_ranks = len(self.choices.columns)
        valid_index = set(self.projects.index)
        self.DecimalChoices = copy.copy(self.choices)

        for student in self.choices.index:
            expanded = []
            for choice in self.choices.loc[student].values:
                expanded.append(choice)
                if pandas.isna(choice):
                    continue
                addition = 1
                while True:
                    sibling = f"{choice}_{addition}"
                    if sibling in valid_index:
                        expanded.append(sibling)
                        addition += 1
                    else:
                        break

            expanded = expanded[:n_ranks]
            while len(expanded) < n_ranks:
                expanded.append(np.nan)
            for i in range(n_ranks):
                self.DecimalChoices.at[student, str(i + 1)] = expanded[i]

    def calc_CostMatrix(self, adjustments=None):
        #adjustments = [{"Baum": 0.1}]
        self.CostMatrix = pandas.DataFrame(index=self.DecimalChoices.index, columns=self.projects.index)
        #its a cost matrix so we want to price the choices low and the projects not chosen very highly
        HighCost = 10000
        self.CostMatrix[:] = HighCost

        valid_projects = set(self.projects.index)

        for student in self.DecimalChoices.index:
            invalid_count = 0
            valid_choices = []
            for choice in self.DecimalChoices.loc[student].values:
                if pandas.isna(choice):
                    invalid_count += 1
                elif choice in valid_projects:
                    valid_choices.append(choice)
                else:
                    invalid_count += 1

            for cost, choice in enumerate(valid_choices):
                self.CostMatrix.at[student, choice] = (cost**2) + invalid_count

    def Hungarian(self):
        row_ind, col_ind = linear_sum_assignment(self.CostMatrix)

        self.Result = pandas.DataFrame(index=self.choices.index, columns=["Project", "Choice #"])

        for i in range(len(row_ind)):
            student = self.CostMatrix.index[row_ind[i]]
            project_assigned = self.CostMatrix.columns[col_ind[i]]
            base_project = project_assigned.split("_")[0]
            self.Result.at[student, "Project"] = base_project

            student_choices = [c for c in self.choices.loc[student].values if not pandas.isna(c)]
            try:
                self.Result.at[student, "Choice #"] = student_choices.index(base_project) + 1
            except ValueError:
                self.Result.at[student, "Choice #"] = -1

        print("Assignment complete.")

    def summary(self):
        print("="*50)
        print("Result summary:")
        print("Number of students that got their first choice:", (self.Result["Choice #"] == 1).sum())
        print("Number of students that got their second choice:", (self.Result["Choice #"] == 2).sum())
        print("Number of students that got their third choice:", (self.Result["Choice #"] == 3).sum())
        print("Number of students that got their fourth choice:", (self.Result["Choice #"] == 4).sum())
        print("Number of students that got their fifth choice:", (self.Result["Choice #"] == 5).sum())
        print("Lowest choice:", self.Result["Choice #"].max())

        print("Average choice that each student got:", round(self.Result[self.Result["Choice #"] != -1]["Choice #"].values.astype(np.float64).mean(), 2))
        print("Did any student get assigned a project that wasn't on their list of choices:", (self.Result["Choice #"] == -1).any())

        self.categories = {"green": {"assigned":0, "cap":12},
                           "Catalysis (blue)": {"assigned":0, "cap":22},
                           "BioNano (orange)": {"assigned":0, "cap":35},
                           "MCCB (purple)": {"assigned":0, "cap":28},
                           "red": {"assigned":0, "cap":9}}

        for project_choice in self.Result["Project"]:
            leading_digits = re.match(r"\d+", str(project_choice))
            if not leading_digits:
                print(f"Warning: could not categorise project ID: {project_choice}")
                continue
            project_num = int(leading_digits.group())
            if project_num <= 12:
                self.categories["green"]["assigned"]  += 1
            elif project_num<= 33:
                self.categories["Catalysis (blue)"]["assigned"]  += 1
            elif project_num<= 57:
                self.categories["BioNano (orange)"]["assigned"]  += 1
            elif project_num<= 82:
                self.categories["MCCB (purple)"]["assigned"]  += 1
            else:
                self.categories["red"]["assigned"]  += 1
        for colour in self.categories:
            print(f"{colour} had capacity for:", self.categories[colour]["cap"],
                  "and got assigned", self.categories[colour]["assigned"],
                  "\t ratio:", self.categories[colour]["assigned"]/self.categories[colour]["cap"], ": 1")

        #by-sectiongroup-ratio (MCCB offered 15 : 80, BIONANO X:N etc)
        print("="*50)


    def capacity_check(self):
        print("Checking that each supervisor has atleast 1 student")
        print("Checking that no supervisor is oversubscribed")
        supervisors = np.unique(self.projects["Section"])
        print(supervisors)
        for supervisor in supervisors:
            subset = self.projects[self.projects["Section"] == supervisor]
            capacity = subset["Capacity"].sum()

            base_projects = {p.split("_")[0] for p in subset.index}
            assigned_students = self.Result[self.Result["Project"].isin(base_projects)].shape[0]

            print(supervisor, capacity, assigned_students)
            assert capacity >= assigned_students

            if assigned_students == 0:
                print(list(base_projects))
                chosen_anywhere = self.DecimalChoices.isin(base_projects).any().any()
                print("Nobody chose any of these projects:", not chosen_anywhere)
        return None


    def save(self, output_path=None):
        output_path = os.path.abspath(output_path or "./Project-assignment.csv")
        self.Result.to_csv(output_path)
        print(f"\nResults saved to: {output_path}")

    def __init__(self, projects, choices):
        self.load_projects(projects)
        self.load_choices(choices)
        self.calc_CostMatrix()

def _clean_path(text):
    """Strip whitespace, a stray BOM, and any quotes Windows adds when you paste a path."""
    return text.strip().lstrip("﻿").strip('"').strip("'")


if __name__ == "__main__":
    print("Student Project Assignment - Hungarian algorithm")
    print("Enter the path to each file below (a .csv or .xlsx file works).\n")

    projects = _clean_path(input("Projects file: "))
    student_selections = _clean_path(input("Student choices file: "))
    print()

    try:
        spa = StudentProjectAssignment(projects, student_selections)
        spa.Hungarian()
        spa.summary()
        spa.save()
    except Exception as exc:
        print(f"\nSomething went wrong: {exc}")

    try:
        input("\nPress Enter to close this window...")
    except EOFError:
        pass
